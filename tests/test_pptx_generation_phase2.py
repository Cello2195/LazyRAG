from __future__ import annotations

import os
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
ALGO = ROOT / 'algorithm'
if str(ALGO) not in sys.path:
    sys.path.insert(0, str(ALGO))

from chat.pptx.schema import (  # noqa: E402
    normalize_deck_schema,
    repair_deck_schema_for_pptx,
    validate_deck_schema,
)
from chat.tools import pptx as pptx_tools  # noqa: E402


def test_normalize_handles_missing_fields():
    schema = {'slides': [{'bullets': 'A\nB\nC'}]}
    normalized = normalize_deck_schema(schema)

    assert normalized['title']
    assert normalized['slides']
    assert normalized['slides'][0]['title']
    assert normalized['slides'][0]['type'] in {
        'cover',
        'content_bullets',
        'summary',
        'toc',
        'table',
        'two_column',
        'comparison',
        'section_divider',
        'references',
    }


def test_validate_returns_normalized_schema():
    schema = {'title': 'Demo', 'slides': [{'type': 'table', 'headers': ['A'], 'rows': [['1']]}]}
    result = validate_deck_schema(schema)

    assert result['success'] is True
    assert 'normalized_schema' in result
    assert isinstance(result['normalized_schema'], dict)


def test_phase2_generation_parse_qa_repair_and_artifact(tmp_path, monkeypatch):
    monkeypatch.setenv('LAZYRAG_UPLOAD_ROOT', str(tmp_path / 'uploads'))

    schema = {
        'title': 'Phase2 Unit Test',
        'language': 'zh',
        'theme': 'business',
        'slides': [
            {'type': 'cover', 'title': '封面', 'subtitle': '测试'},
            {
                'type': 'content_bullets',
                'title': '内容',
                'bullets': [
                    '这是一个比较长的测试条目，用于验证文本长度控制与解析稳定性。',
                    'TODO placeholder text should be repaired',
                    '第三条内容',
                    '第四条内容',
                    '第五条内容',
                    '第六条内容',
                    '第七条内容',
                ],
            },
            {
                'type': 'table',
                'title': '表格',
                'headers': ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H'],
                'rows': [[str(i), str(i + 1), str(i + 2), str(i + 3), str(i + 4), str(i + 5), str(i + 6), str(i + 7)] for i in range(15)],
            },
            {'type': 'summary', 'title': '总结', 'takeaways': ['x', 'y', 'z']},
        ],
        'sources': ['https://example.com/ref1'],
    }

    created = pptx_tools.pptx_create_from_schema(schema, filename='phase2_test.pptx')
    assert created['success'] is True
    file_path = Path(created['file_path']).resolve()
    assert file_path.exists()

    parsed = pptx_tools.pptx_parse(str(file_path))
    assert parsed['success'] is True
    assert int(parsed['slide_count']) >= 4

    qa = pptx_tools.pptx_qa(str(file_path), deck_schema=schema)
    assert qa['success'] is True
    assert 'summary' in qa
    assert isinstance(qa['issues'], list)
    assert isinstance(qa['warnings'], list)

    repaired = repair_deck_schema_for_pptx(schema, qa)
    assert repaired['success'] is True
    assert isinstance(repaired['repaired_schema'], dict)

    saved = pptx_tools.artifact_save(str(file_path), kind='pptx', filename='phase2_saved.pptx')
    assert saved['success'] is True
    saved_path = Path(saved['file_path']).resolve()
    assert saved_path.exists()
    assert saved['sha256']
    assert saved['created_at']


def test_qa_detects_placeholder(tmp_path, monkeypatch):
    monkeypatch.setenv('LAZYRAG_UPLOAD_ROOT', str(tmp_path / 'uploads'))

    schema = {
        'title': 'Placeholder QA Test',
        'slides': [
            {'type': 'cover', 'title': '封面'},
            {'type': 'content_bullets', 'title': '页2', 'bullets': ['TODO this must be flagged']},
        ],
    }
    created = pptx_tools.pptx_create_from_schema(schema, filename='placeholder_qa.pptx')

    # Inject placeholder residue directly into generated PPTX to verify QA detection.
    prs = Presentation(created['file_path'])
    slide = prs.slides[1]
    marker = slide.shapes.add_textbox(Inches(0.8), Inches(5.8), Inches(6.5), Inches(0.4))
    marker.text_frame.text = 'TODO placeholder marker'
    prs.save(created['file_path'])

    qa = pptx_tools.pptx_qa(created['file_path'], deck_schema=schema)

    placeholder_issues = [issue for issue in qa.get('issues', []) if issue.get('code') == 'placeholder_residue']
    assert placeholder_issues
