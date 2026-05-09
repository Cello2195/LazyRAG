from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, Iterable, List

from pptx import Presentation
from pptx.util import Inches


def _resolve_images(paths: Iterable[str]) -> List[Path]:
    images: List[Path] = []
    for raw in paths or []:
        p = Path(raw).expanduser().resolve()
        if p.exists() and p.is_file():
            images.append(p)
    return images


def create_pptx_from_slide_images(
    image_paths: Iterable[str],
    output_path: str | Path,
    *,
    slide_width: float = 13.333,
    slide_height: float = 7.5,
) -> Dict[str, Any]:
    images = _resolve_images(image_paths)
    if not images:
        return {
            'success': False,
            'error_message': 'no valid screenshot image paths supplied',
            'slide_count': 0,
            'image_paths': [],
        }

    out = Path(output_path).expanduser().resolve()
    out.parent.mkdir(parents=True, exist_ok=True)

    prs = Presentation()
    prs.slide_width = Inches(slide_width)
    prs.slide_height = Inches(slide_height)
    blank = prs.slide_layouts[6]

    for image in images:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(image), Inches(0), Inches(0), width=prs.slide_width, height=prs.slide_height)

    prs.save(str(out))
    size_bytes = out.stat().st_size if out.exists() else 0
    return {
        'success': True,
        'file_path': str(out),
        'filename': out.name,
        'slide_count': len(images),
        'size_bytes': size_bytes,
        'image_paths': [str(p) for p in images],
        'editable': False,
        'generation_mode': 'html_screenshot_image_pptx',
        'note': 'visual_pptx is image-based; slide elements are not fully editable.',
    }
